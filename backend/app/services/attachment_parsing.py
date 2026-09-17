from __future__ import annotations

import asyncio
import csv
import json
import logging
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.models.conversation import Attachment
from app.models.enums import AttachmentParseStatus
from app.repositories.attachments import update_parse_result
from app.services.attachments import resolve_attachment_path

logger = logging.getLogger(__name__)


class AttachmentParsingError(Exception):
    """A safe, user-facing attachment parsing error."""


@dataclass(frozen=True)
class ParsingLimits:
    max_rows: int
    max_columns: int
    max_sheets: int
    max_text_chars: int
    max_cell_chars: int
    max_json_depth: int

    @classmethod
    def from_settings(cls, settings: Settings) -> ParsingLimits:
        return cls(
            max_rows=settings.attachment_parse_max_rows,
            max_columns=settings.attachment_parse_max_columns,
            max_sheets=settings.attachment_parse_max_sheets,
            max_text_chars=settings.attachment_parse_max_text_chars,
            max_cell_chars=settings.attachment_parse_max_cell_chars,
            max_json_depth=settings.attachment_parse_max_json_depth,
        )


def _read_text(path: Path, max_chars: int) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open(encoding=encoding) as source:
                content = source.read(max_chars + 1)
            if len(content) > max_chars:
                raise AttachmentParsingError(f"文本内容超过 {max_chars} 个字符的解析上限")
            return content, encoding
        except UnicodeDecodeError:
            continue
    raise AttachmentParsingError("文件编码无法识别，请使用 UTF-8 或 GB18030 编码")


def _normalized_headers(values: list[str]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = value.strip() or f"column_{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _validate_cell(value: Any, limits: ParsingLimits) -> Any:
    if value is None or isinstance(value, str | int | float | bool):
        normalized = value
    elif isinstance(value, datetime | date | time):
        normalized = value.isoformat()
    elif isinstance(value, timedelta):
        normalized = str(value)
    elif isinstance(value, Decimal):
        normalized = str(value)
    else:
        normalized = str(value)
    if isinstance(normalized, str) and len(normalized) > limits.max_cell_chars:
        raise AttachmentParsingError(
            f"单元格内容超过 {limits.max_cell_chars} 个字符的解析上限"
        )
    return normalized


def _table_payload(
    raw_rows: list[list[Any]],
    *,
    source_format: str,
    limits: ParsingLimits,
    name: str | None = None,
) -> dict[str, Any]:
    if not raw_rows:
        columns: list[str] = []
        rows: list[dict[str, Any]] = []
    else:
        width = max(len(row) for row in raw_rows)
        if width > limits.max_columns:
            raise AttachmentParsingError(f"数据列数超过 {limits.max_columns} 列的解析上限")
        header_values = [str(value or "") for value in raw_rows[0]]
        header_values.extend("" for _ in range(width - len(header_values)))
        columns = _normalized_headers(header_values)
        rows = []
        for raw_row in raw_rows[1:]:
            padded = [*raw_row, *([None] * (width - len(raw_row)))]
            rows.append(
                {
                    column: _validate_cell(value, limits)
                    for column, value in zip(columns, padded, strict=True)
                }
            )

    payload: dict[str, Any] = {
        "kind": "table",
        "source_format": source_format,
        "columns": columns,
        "rows": rows,
        "row_count": len(rows),
    }
    if name is not None:
        payload["name"] = name
    return payload


def _parse_csv(path: Path, limits: ParsingLimits) -> dict[str, Any]:
    last_decode_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            raw_rows: list[list[Any]] = []
            with path.open(encoding=encoding, newline="") as source:
                for row_number, row in enumerate(csv.reader(source), start=1):
                    if row_number > limits.max_rows + 1:
                        raise AttachmentParsingError(
                            f"CSV 数据行超过 {limits.max_rows} 行的解析上限"
                        )
                    if len(row) > limits.max_columns:
                        raise AttachmentParsingError(
                            f"CSV 数据列超过 {limits.max_columns} 列的解析上限"
                        )
                    raw_rows.append(row)
            payload = _table_payload(
                raw_rows,
                source_format="csv",
                limits=limits,
            )
            payload["encoding"] = encoding
            return payload
        except UnicodeDecodeError as error:
            last_decode_error = error
    raise AttachmentParsingError(
        "CSV 编码无法识别，请使用 UTF-8 或 GB18030 编码"
    ) from last_decode_error


def _parse_xlsx(path: Path, limits: ParsingLimits) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as error:
        raise AttachmentParsingError("XLSX 文件损坏或格式无效") from error

    try:
        if len(workbook.worksheets) > limits.max_sheets:
            raise AttachmentParsingError(
                f"XLSX 工作表超过 {limits.max_sheets} 个的解析上限"
            )
        sheets: list[dict[str, Any]] = []
        total_rows = 0
        for worksheet in workbook.worksheets:
            raw_rows: list[list[Any]] = []
            for row in worksheet.iter_rows(values_only=True):
                values = list(row)
                while values and values[-1] is None:
                    values.pop()
                if not values and not raw_rows:
                    continue
                raw_rows.append(values)
                if len(raw_rows) > limits.max_rows + 1:
                    raise AttachmentParsingError(
                        f"工作表“{worksheet.title}”超过 {limits.max_rows} 行的解析上限"
                    )
            table = _table_payload(
                raw_rows,
                source_format="xlsx",
                limits=limits,
                name=worksheet.title,
            )
            total_rows += table["row_count"]
            sheets.append(table)
        return {
            "kind": "workbook",
            "source_format": "xlsx",
            "sheets": sheets,
            "sheet_count": len(sheets),
            "row_count": total_rows,
        }
    finally:
        workbook.close()


def _json_depth(value: Any, max_depth: int) -> int:
    stack = [(value, 1)]
    deepest = 0
    while stack:
        current, depth = stack.pop()
        if depth > max_depth:
            raise AttachmentParsingError(f"JSON 嵌套层级超过 {max_depth} 层的解析上限")
        deepest = max(deepest, depth)
        if isinstance(current, dict):
            stack.extend((item, depth + 1) for item in current.values())
        elif isinstance(current, list):
            stack.extend((item, depth + 1) for item in current)
    return deepest


def _parse_json(path: Path, limits: ParsingLimits) -> dict[str, Any]:
    content, encoding = _read_text(path, limits.max_text_chars)
    try:
        data = json.loads(content)
    except json.JSONDecodeError as error:
        raise AttachmentParsingError(
            f"JSON 格式错误：第 {error.lineno} 行第 {error.colno} 列"
        ) from error
    depth = _json_depth(data, limits.max_json_depth)
    return {
        "kind": "json",
        "source_format": "json",
        "encoding": encoding,
        "depth": depth,
        "data": data,
    }


def _parse_txt(path: Path, limits: ParsingLimits) -> dict[str, Any]:
    content, encoding = _read_text(path, limits.max_text_chars)
    return {
        "kind": "text",
        "source_format": "txt",
        "encoding": encoding,
        "character_count": len(content),
        "content": content,
    }


def parse_attachment_file(path: Path, limits: ParsingLimits) -> dict[str, Any]:
    parsers = {
        ".csv": _parse_csv,
        ".xlsx": _parse_xlsx,
        ".json": _parse_json,
        ".txt": _parse_txt,
    }
    parser = parsers.get(path.suffix.lower())
    if parser is None:
        raise AttachmentParsingError("该文件类型不支持解析")
    return parser(path, limits)


async def parse_attachment_record(
    session: AsyncSession,
    attachment: Attachment,
    settings: Settings,
) -> Attachment:
    await update_parse_result(
        session,
        attachment,
        status=AttachmentParseStatus.PARSING,
    )
    path = resolve_attachment_path(settings.storage_path, attachment.file_path)
    try:
        if not path.is_file():
            raise AttachmentParsingError("附件文件不存在")
        parsed = await asyncio.to_thread(
            parse_attachment_file,
            path,
            ParsingLimits.from_settings(settings),
        )
    except AttachmentParsingError as error:
        return await update_parse_result(
            session,
            attachment,
            status=AttachmentParseStatus.FAILED,
            error=str(error),
        )
    except Exception:
        logger.exception("Unexpected attachment parsing error attachment_id=%s", attachment.id)
        return await update_parse_result(
            session,
            attachment,
            status=AttachmentParseStatus.FAILED,
            error="文件解析失败，请检查文件内容后重试",
        )

    return await update_parse_result(
        session,
        attachment,
        status=AttachmentParseStatus.SUCCESS,
        content=parsed,
    )
